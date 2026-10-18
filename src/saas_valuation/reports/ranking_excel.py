"""Writes the ranking workbook: final ranking, comps regression (with chart), methodology."""

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.styles import Font, PatternFill

from saas_valuation.reports.comps_excel import (BUCKET_NAMES, MULTIPLE, PERCENT, PRICE, _header,
                                                _text_sheet, _title)
from saas_valuation.reports.dcf_excel import ABOVE, BELOW, _cell

METRIC_LABELS = {"revenue_growth": "Revenue Growth", "rule_of_40": "Rule of 40", "fcf_margin": "FCF Margin"}
NEUTRAL = PatternFill("solid", fgColor="FFF2CC")
VERDICT_FILL = {"Undervalued": ABOVE, "Overvalued": BELOW, "In line": NEUTRAL}

COLUMNS = [
    ("rank", "Rank", "0", 6), ("name", "Company", None, 28), ("bucket", "Profile", None, 22),
    ("verdict", "Verdict", None, 13), ("score", "Score (0-100)", "0", 9),
    ("price", "Share Price", PRICE, 11), ("dcf_price", "DCF Price", PRICE, 11),
    ("comps_price", "Comps Price", PRICE, 11), ("blended_price", "Blended Price", PRICE, 11),
    ("dcf_upside", "DCF Upside", PERCENT, 10), ("comps_upside", "Comps Upside", PERCENT, 10),
    ("blended_upside", "Blended Upside", PERCENT, 10), ("dcf_score", "DCF Score", "0", 8),
    ("comps_score", "Comps Score", "0", 8), ("signals_agree", "Signals Agree?", None, 10),
    ("ev_revenue", "EV / Revenue", MULTIPLE, 10), ("fair_ev_revenue", "Fair EV / Revenue", MULTIPLE, 10),
]


def _ranking_sheet(ws, ranking: pd.DataFrame, subtitle: str) -> None:
    _title(ws, "SaaS Valuation Ranking", subtitle)
    _header(ws, 4, ["Ticker"] + [h for _, h, _, _ in COLUMNS])
    for r, (ticker, rec) in enumerate(ranking.iterrows(), start=5):
        _cell(ws, r, 1, ticker, bold=True)
        for c, (col, _, fmt, _) in enumerate(COLUMNS, start=2):
            value = rec[col]
            fill = None
            if col == "bucket":
                value = BUCKET_NAMES.get(value, value)
            elif col == "signals_agree":
                value = "Yes" if value else "No"
            elif col == "verdict":
                fill = VERDICT_FILL.get(value)
            elif col.endswith("upside") and pd.notna(value):
                fill = ABOVE if value > 0 else BELOW
            _cell(ws, r, c, value, fmt, bold=col in ("verdict", "score"), fill=fill)
    ws.freeze_panes = "C5"
    ws.column_dimensions["A"].width = 9
    for c, (_, _, _, width) in enumerate(COLUMNS, start=2):
        ws.column_dimensions[ws.cell(row=4, column=c).column_letter].width = width


def _regression_sheet(ws, ranking: pd.DataFrame, reg: dict) -> None:
    label = METRIC_LABELS[reg["regressor"]]
    _title(ws, f"Comps regression: EV / Revenue vs {label}",
           f"ln(EV/Revenue) = {reg['intercept']:.2f} + {reg['slope']:.2f} x {label}   ·   "
           f"R² = {reg['r_squared']:.2f}   ·   n = {reg['n']}   ·   "
           f"+10 pts of {label} -> multiple {reg['multiple_change_per_10pts']:+.0%}")

    data = ranking.dropna(subset=["regressor", "ev_revenue"]).sort_values("regressor")
    _header(ws, 4, ["Ticker", label, "EV / Revenue", "Fair EV / Revenue", "Fitted line",
                    "Premium / (Discount) to Fair"])
    for r, (ticker, rec) in enumerate(data.iterrows(), start=5):
        fitted = np.exp(reg["intercept"] + reg["slope"] * rec["regressor"])
        premium = rec["ev_revenue"] / rec["fair_ev_revenue"] - 1
        _cell(ws, r, 1, ticker, bold=True)
        _cell(ws, r, 2, rec["regressor"], PERCENT)
        _cell(ws, r, 3, rec["ev_revenue"], MULTIPLE)
        _cell(ws, r, 4, rec["fair_ev_revenue"], MULTIPLE)
        _cell(ws, r, 5, fitted, MULTIPLE)
        _cell(ws, r, 6, premium, PERCENT, fill=BELOW if premium > 0 else ABOVE)
    last = 4 + len(data)
    for col, width in zip("ABCDEF", (9, 13, 13, 15, 12, 18)):
        ws.column_dimensions[col].width = width
    ws.cell(row=last + 2, column=1,
            value="Fair EV/Revenue uses a line fitted on the other companies (leave-one-out), so it "
                  "differs slightly from the fitted line, which uses all of them.").font = Font(italic=True)

    row = last + 4
    ws.cell(row=row, column=1, value="Which metric explains the multiples best?").font = Font(bold=True, size=12)
    _header(ws, row + 1, ["Metric", "R²", "Used"])
    for i, (metric, r2) in enumerate(sorted(reg["candidates"].items(), key=lambda kv: -kv[1]), start=row + 2):
        _cell(ws, i, 1, METRIC_LABELS[metric], bold=True)
        _cell(ws, i, 2, r2, "0.00")
        _cell(ws, i, 3, "Yes" if metric == reg["regressor"] else "")
    ws.cell(row=i + 1, column=1,
            value="R² = share of the variation in ln(EV/Revenue) the metric explains (1.0 = all of it).").font = Font(italic=True)

    chart = ScatterChart()
    chart.title = f"EV / Revenue vs {label}"
    chart.style = 13
    chart.x_axis.title = label
    chart.y_axis.title = "EV / Revenue (log scale)"
    chart.y_axis.scaling.logBase = 10
    chart.x_axis.number_format = "0%"
    chart.height, chart.width = 12, 20
    x_ref = Reference(ws, min_col=2, min_row=5, max_row=last)
    points = Series(Reference(ws, min_col=3, min_row=4, max_row=last), x_ref, title_from_data=True)
    points.marker.symbol, points.marker.size = "circle", 8
    points.graphicalProperties.line.noFill = True
    line = Series(Reference(ws, min_col=5, min_row=4, max_row=last), x_ref, title_from_data=True)
    line.marker.symbol = "none"
    chart.series += [points, line]
    ws.add_chart(chart, "H4")


def write_ranking_workbook(path: Path, ranking: pd.DataFrame, reg: dict, cfg: dict) -> None:
    r = cfg["ranking"]
    subtitle = (f"Score = {r['dcf_upside_weight']:.0%} DCF percentile + {r['comps_weight']:.0%} comps "
                f"percentile · Undervalued ≥ {r['undervalued_above']}, Overvalued ≤ {r['overvalued_below']} · "
                "relative to this 18-company universe")
    wb = Workbook()
    wb.active.title = "Ranking"
    _ranking_sheet(wb.active, ranking, subtitle)
    _regression_sheet(wb.create_sheet("Comps Regression"), ranking, reg)
    label = METRIC_LABELS[r["comps_regressor"]]
    _text_sheet(wb.create_sheet("Methodology"), "How the ranking works", [
        "1. DCF signal",
        "  DCF upside = DCF implied price / current price - 1 (Stage 4: two-stage DCF, SBC expensed).",
        "",
        "2. Comps signal",
        "  Raw multiples aren't comparable: faster-growing, more profitable companies deserve higher ones.",
        f"  So ln(EV/Revenue) is regressed on {label} across the universe. Each company's 'fair' multiple",
        "  comes from a line fitted on the other companies (leave-one-out), so it can't set its own benchmark.",
        "  Comps price = (fair multiple x revenue + net cash) / shares. Comps upside = comps price / price - 1.",
        "",
        "3. Combining",
        "  Each upside is converted to a percentile rank (0-100) within the universe, then weighted:",
        f"  score = {r['dcf_upside_weight']} x DCF percentile + {r['comps_weight']} x comps percentile.",
        "  Percentiles stop one extreme result from dominating; they measure order, not size.",
        "  The blended price is the same weighted average applied to the two implied prices.",
        "",
        "4. Reading the result",
        "  Verdicts are relative: 'Undervalued' means cheaper than peers on these two views, not a buy signal.",
        "  Signals agree = DCF and comps point the same way. Agreement means higher conviction;",
        "  disagreement usually means the market expects growth beyond what the DCF assumes.",
        "  All weights and thresholds are set in config/config.yaml.",
    ])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
