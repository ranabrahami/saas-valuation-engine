"""Writes the DCF workbook: a summary sheet plus one model sheet per company."""

from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from saas_valuation.reports.comps_excel import (BUCKET_NAMES, HEADER_FILL, HEADER_FONT, LABEL_FONT,
                                                MILLIONS, MULTIPLE, PERCENT, PRICE, STATS_FILL,
                                                THIN_BOTTOM, _header, _title)

ABOVE = PatternFill("solid", fgColor="D9EAD3")   # implied price above current price
BELOW = PatternFill("solid", fgColor="F4CCCC")
BASE = Font(bold=True)
DECIMAL = "0.00"

SUMMARY_COLUMNS = [
    ("name", "Company", None, 30), ("bucket", "Profile", None, 24),
    ("price", "Share Price", PRICE, 12), ("implied_price", "DCF Price", PRICE, 12),
    ("upside", "Upside / (Downside)", PERCENT, 13), ("wacc", "WACC", PERCENT, 9),
    ("cost_of_equity", "Cost of Equity", PERCENT, 10), ("beta", "Beta (adj.)", DECIMAL, 9),
    ("start_growth", "Starting Growth", PERCENT, 10),
    ("market_implied_growth", "Growth Implied by Price", PERCENT, 12),
    ("year5_growth", "Year-5 Growth", PERCENT, 10),
    ("start_fcf_margin", "Starting FCF Margin", PERCENT, 11),
    ("target_fcf_margin", "Mature FCF Margin", PERCENT, 11), ("tv_pct_ev", "TV % of EV", PERCENT, 10),
    ("implied_terminal_ev_revenue", "Implied Terminal EV/Rev", MULTIPLE, 12),
    ("market_ev_revenue", "Market EV/Rev Today", MULTIPLE, 12), ("notes", "Notes", None, 60),
]


def _cell(ws, row, col, value, fmt=None, bold=False, fill=None):
    cell = ws.cell(row=row, column=col)
    if value is None or (not isinstance(value, str) and pd.isna(value)):
        cell.value, cell.font, cell.alignment = "NM", LABEL_FONT, Alignment(horizontal="right")
    else:
        cell.value = value
        if fmt:
            cell.number_format = fmt
        if bold:
            cell.font = BASE
    if fill:
        cell.fill = fill
    return cell


def _summary_sheet(ws, summary: pd.DataFrame, subtitle: str) -> None:
    _title(ws, "DCF Valuation Summary", subtitle)
    _header(ws, 4, ["Ticker"] + [h for _, h, _, _ in SUMMARY_COLUMNS])
    ordered = summary.sort_values("upside", ascending=False, na_position="last")
    for r, (ticker, rec) in enumerate(ordered.iterrows(), start=5):
        _cell(ws, r, 1, ticker, bold=True)
        for c, (col, _, fmt, _) in enumerate(SUMMARY_COLUMNS, start=2):
            value = BUCKET_NAMES.get(rec[col], rec[col]) if col == "bucket" else rec[col]
            if col == "notes" and (not isinstance(value, str) or not value):
                value = ""
            fill = None
            if col == "upside" and pd.notna(value):
                fill = ABOVE if value > 0 else BELOW
            _cell(ws, r, c, value, fmt, fill=fill)
    ws.freeze_panes = "B5"
    ws.column_dimensions["A"].width = 10
    for c, (_, _, _, width) in enumerate(SUMMARY_COLUMNS, start=2):
        ws.column_dimensions[ws.cell(row=4, column=c).column_letter].width = width


def _company_sheet(ws, ticker: str, rec: pd.Series, result: dict, cfg: dict, subtitle: str) -> None:
    d = cfg["dcf"]
    _title(ws, f"{ticker} · {rec['name']} — DCF", subtitle)
    ws.column_dimensions["A"].width = 34
    for col in "BCDEFGHIJKL":
        ws.column_dimensions[col].width = 12

    # --- Assumptions ---
    row = 4
    ws.cell(row=row, column=1, value="Assumptions").font = Font(bold=True, size=12)
    margin_label = "FCF margin after SBC" if d["sbc_treatment"] == "expense" else "FCF margin"
    assumptions = [
        ("Beta", rec["beta"], DECIMAL, rec["beta_basis"]),
        ("Risk-free rate", d["risk_free_rate"], PERCENT, "10-year Treasury"),
        ("Equity risk premium", d["equity_risk_premium"], PERCENT, ""),
        ("Cost of equity (CAPM)", rec["cost_of_equity"], PERCENT, "risk-free + beta x ERP"),
        ("Debt / (debt + equity)", rec["debt_weight"], PERCENT, "market-value weights"),
        ("WACC", rec["wacc"], PERCENT, ""),
        ("Terminal growth", d["terminal_growth_rate"], PERCENT, "Gordon growth"),
        ("Starting revenue growth", rec["start_growth"], PERCENT, f"CY{result['base_year']} actual"),
        ("Year-5 revenue growth", rec["year5_growth"], PERCENT, "end of explicit stage"),
        (f"Starting {margin_label}", rec["start_fcf_margin"], PERCENT, f"CY{result['base_year']} actual"),
        (f"Mature {margin_label}", rec["target_fcf_margin"], PERCENT,
         f"reached in year {d['projection_years'] + d['fade_years']}, straight-line path"),
    ]
    for label, value, fmt, note in assumptions:
        row += 1
        ws.cell(row=row, column=1, value=label)
        _cell(ws, row, 2, value, fmt)
        ws.cell(row=row, column=3, value=note).font = LABEL_FONT

    # --- Projection ---
    row += 2
    ws.cell(row=row, column=1, value="Projection ($M)").font = Font(bold=True, size=12)
    proj = result["projection"]
    row += 1
    _header(ws, row, [""] + [f"CY{result['base_year']}A"] + [f"CY{y}E" for y in proj.index])
    row += 1
    ws.cell(row=row, column=1, value="Stage").font = LABEL_FONT
    ws.cell(row=row, column=2, value="actual").font = LABEL_FONT
    for c, stage in enumerate(proj["stage"], start=3):
        cell = ws.cell(row=row, column=c, value=stage)
        cell.font, cell.alignment = LABEL_FONT, Alignment(horizontal="center")
        if stage == "fade":
            cell.fill = STATS_FILL
    lines = [
        ("Revenue", [result["base_revenue"] / 1e6] + list(proj["revenue"] / 1e6), MILLIONS),
        ("Growth", [None] + list(proj["growth"]), PERCENT),
        (margin_label.capitalize(), [rec["start_fcf_margin"]] + list(proj["fcf_margin"]), PERCENT),
        ("Free cash flow", [result["base_revenue"] * rec["start_fcf_margin"] / 1e6] + list(proj["fcf"] / 1e6), MILLIONS),
        ("Discount period (years)", [None] + list(proj["period"]), DECIMAL),
        ("Discount factor", [None] + list(proj["discount_factor"]), "0.000"),
        ("PV of free cash flow", [None] + list(proj["pv_fcf"] / 1e6), MILLIONS),
    ]
    for label, values, fmt in lines:
        row += 1
        ws.cell(row=row, column=1, value=label)
        for c, v in enumerate(values, start=2):
            if v is not None:
                _cell(ws, row, c, v, fmt, bold=label in ("Free cash flow", "PV of free cash flow"))

    # --- Valuation bridge ---
    row += 2
    ws.cell(row=row, column=1, value="Valuation ($M except per share)").font = Font(bold=True, size=12)
    first_year = proj.index[0]
    bridge = [
        ("Sum of PV of free cash flow", rec["sum_pv_fcf"] / 1e6, MILLIONS, False),
        ("PV of terminal value", rec["pv_terminal_value"] / 1e6, MILLIONS, False),
        ("Enterprise value", rec["enterprise_value"] / 1e6, MILLIONS, True),
        ("+ Net cash (cash + investments - debt)", rec["net_cash"] / 1e6, MILLIONS, False),
        (f"Equity value at {first_year}-01-01", (rec["enterprise_value"] + rec["net_cash"]) / 1e6, MILLIONS, False),
        ("x Roll-forward to valuation date", rec["roll_forward_factor"], "0.000", False),
        ("Equity value today", rec["equity_value"] / 1e6, MILLIONS, True),
        ("÷ Shares outstanding (M)", rec["shares"] / 1e6, "#,##0.0", False),
        ("Implied share price", rec["implied_price"], PRICE, True),
        ("Current share price", rec["price"], PRICE, False),
        ("Upside / (downside)", rec["upside"], PERCENT, True),
        ("Starting growth implied by today's price", rec["market_implied_growth"], PERCENT, False),
        ("Terminal value as % of EV", rec["tv_pct_ev"], PERCENT, False),
        ("Implied terminal EV / revenue", rec["implied_terminal_ev_revenue"], MULTIPLE, False),
    ]
    for label, value, fmt, bold in bridge:
        row += 1
        ws.cell(row=row, column=1, value=label).font = Font(bold=bold)
        cell = _cell(ws, row, 2, value, fmt, bold=bold)
        if bold:
            cell.border = THIN_BOTTOM
        if label.startswith("Upside") and pd.notna(value):
            cell.fill = ABOVE if value > 0 else BELOW

    # --- Sensitivity ---
    row += 2
    ws.cell(row=row, column=1,
            value="Sensitivity: implied share price (WACC down, terminal growth across)").font = Font(bold=True, size=12)
    grid = result["sensitivity"]
    row += 1
    ws.cell(row=row, column=1, value="WACC \\ g").font = HEADER_FONT
    ws.cell(row=row, column=1).fill = HEADER_FILL
    for c, g in enumerate(grid.columns, start=2):
        _cell(ws, row, c, g, PERCENT).font = HEADER_FONT
        ws.cell(row=row, column=c).fill = HEADER_FILL
    base_rate, base_g = rec["wacc"], d["terminal_growth_rate"]
    for r_value, values in grid.iterrows():
        row += 1
        label = _cell(ws, row, 1, r_value, PERCENT, bold=True)
        label.fill = STATS_FILL
        for c, (g, price) in enumerate(values.items(), start=2):
            fill = None if pd.isna(price) else (ABOVE if price > rec["price"] else BELOW)
            cell = _cell(ws, row, c, price, PRICE, fill=fill)
            if abs(r_value - base_rate) < 1e-9 and abs(g - base_g) < 1e-9:
                cell.font = Font(bold=True, underline="single")
    row += 1
    ws.cell(row=row, column=1, value=f"Green = above today's price (${rec['price']:,.2f}); "
                                     "red = below; underlined = base case; NM = WACC ≤ growth or "
                                     "negative equity.").font = LABEL_FONT

    if isinstance(rec["notes"], str) and rec["notes"]:
        row += 2
        ws.cell(row=row, column=1, value=f"Notes: {rec['notes']}").font = LABEL_FONT


def write_dcf_workbook(path: Path, summary: pd.DataFrame, results: dict, cfg: dict,
                       valuation_date: date) -> None:
    d = cfg["dcf"]
    subtitle = (f"Valuation date {valuation_date} · $ in millions except per share · "
                f"SBC treated as {'an expense' if d['sbc_treatment'] == 'expense' else 'non-cash (added back)'} · "
                "assumptions in config/config.yaml")
    wb = Workbook()
    _summary_sheet(wb.active, summary, subtitle)
    wb.active.title = "Summary"
    for ticker, rec in summary.sort_values("upside", ascending=False, na_position="last").iterrows():
        _company_sheet(wb.create_sheet(ticker), ticker, rec, results[ticker], cfg, subtitle)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
