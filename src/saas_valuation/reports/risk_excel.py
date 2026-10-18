"""Writes the risk summaries workbook: AI disruption vs valuation, top risks, run log."""

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from saas_valuation.reports.comps_excel import PERCENT, _header, _title
from saas_valuation.reports.dcf_excel import ABOVE, BELOW, _cell
from saas_valuation.reports.ranking_excel import NEUTRAL, VERDICT_FILL

AI_FILL = {"high": BELOW, "medium": NEUTRAL, "low": ABOVE, "none": PatternFill("solid", fgColor="EDEDED")}
WRAP = Alignment(wrap_text=True, vertical="top")
UNVERIFIED = Font(italic=True, color="C00000")


def _sheet(ws, title, subtitle, headers, rows, widths):
    _title(ws, title, subtitle)
    _header(ws, 4, headers)
    for r, values in enumerate(rows, start=5):
        for c, (value, fmt, fill) in enumerate(values, start=1):
            cell = _cell(ws, r, c, value, fmt, bold=c == 1, fill=fill)
            cell.alignment = WRAP
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=4, column=i).column_letter].width = width
    ws.freeze_panes = "B5"


def write_risk_workbook(path: Path, table: pd.DataFrame, results: list[dict]) -> None:
    wb = Workbook()
    ok = [r for r in results if r.get("status") == "ok"]
    by_ticker = {r["ticker"]: r for r in ok}

    rows = []
    for ticker, rec in table.iterrows():
        rows.append([
            (ticker, None, None), (rec.get("rank"), "0", None), (rec.get("name", ""), None, None),
            (rec.get("verdict", ""), None, VERDICT_FILL.get(rec.get("verdict"))),
            (rec.get("blended_upside"), PERCENT, None),
            (rec["ai_level"].upper(), None, AI_FILL.get(rec["ai_level"])),
            (rec["ai_mechanisms"].replace("_", " "), None, None), (rec["ai_assessment"], None, None),
            (f'"{rec["ai_quote"]}"' if rec["ai_quote"] else "", None, None),
            ("Yes" if rec["ai_quote_verified"] else "NOT FOUND", None, None if rec["ai_quote_verified"] else BELOW),
        ])
    _sheet(wb.active, "AI disruption risk vs valuation",
           "From each company's latest 10-K Item 1A, summarized by Claude. Quotes are checked against the filing text.",
           ["Ticker", "Rank", "Company", "Verdict", "Blended Upside", "AI Risk", "Mechanisms", "Assessment",
            "Quote from 10-K", "Quote Verified"], rows, [9, 6, 24, 13, 10, 9, 22, 60, 55, 10])
    wb.active.title = "AI Risk vs Valuation"

    rows = []
    for ticker in table.index:
        for i, risk in enumerate(by_ticker[ticker]["top_risks"], start=1):
            rows.append([
                (ticker, None, None), (i, "0", None), (risk["title"], None, None),
                (risk["category"].replace("_", " "), None, None),
                (risk["severity"], None, {"high": BELOW, "medium": NEUTRAL, "low": ABOVE}[risk["severity"]]),
                ("Yes" if risk["company_specific"] else "Generic", None, None),
                (risk["explanation"], None, None), (f'"{risk["quote"]}"', None, None),
                ("Yes" if risk["quote_verified"] else "NOT FOUND", None, None if risk["quote_verified"] else BELOW),
            ])
    _sheet(wb.create_sheet("Top Risks"), "Most material risks per company",
           "Ranked by likely impact on growth, margins or competitive position, not by order in the filing.",
           ["Ticker", "#", "Risk", "Category", "Severity", "Company-specific?", "Why it matters", "Quote from 10-K",
            "Quote Verified"], rows, [9, 4, 34, 18, 9, 12, 60, 55, 10])

    rows = []
    for r in sorted(results, key=lambda r: r["ticker"]):
        rows.append([(r["ticker"], None, None), (r.get("filing_date"), None, None), (r.get("words"), "#,##0", None),
                     (r.get("model_served"), None, None), (r.get("status"), None, None),
                     (r.get("usage", {}).get("input_tokens"), "#,##0", None),
                     (r.get("usage", {}).get("output_tokens"), "#,##0", None),
                     (r.get("cost_usd"), "$0.00", None), (r.get("url"), None, None)])
    _sheet(wb.create_sheet("Run Log"), "Run log",
           f"Total cost ${sum(r.get('cost_usd', 0) for r in results):.2f} across {len(results)} filings.",
           ["Ticker", "10-K Filed", "Item 1A Words", "Model", "Status", "Input Tokens", "Output Tokens",
            "Cost", "Filing URL"], rows, [9, 11, 12, 18, 10, 12, 12, 9, 80])

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
