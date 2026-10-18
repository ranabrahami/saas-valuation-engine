"""Writes the comps workbook in a banking-style layout.

Sheets:
    Valuation          price, market cap, EV and multiples, with summary stats
    Operating Metrics  growth, margins, Rule of 40, SBC, dilution, with summary stats
    Profile Medians    median metrics per company profile (mature vs high-growth...)
    EV Bridge          market cap -> enterprise value, line by line
    Calendarization    how each company's fiscal years map onto the calendar year
    Stock Splits       every split on record and what it changed
    Overrides          hand-entered values and their sources
    Notes              per-company caveats behind any n/a or NM
    Methodology        metric definitions and the assumptions used
"""

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

MILLIONS = '#,##0'
MULTIPLE = '0.0"x"'
PERCENT = '0.0%'
PRICE = '$#,##0.00'

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
STATS_FILL = PatternFill("solid", fgColor="E7ECF5")
LABEL_FONT = Font(italic=True, color="808080")
THIN_BOTTOM = Border(bottom=Side(style="thin", color="1F3864"))

BUCKET_NAMES = {
    "mature_profitable": "Mature, profitable",
    "growth_profitable": "Growth, profitable",
    "high_growth_unprofitable": "High growth, unprofitable",
    "maturing": "Maturing",
}

# column -> (header, number format, divide by 1e6?)
COLUMNS = {
    "name": ("Company", None, False),
    "bucket": ("Profile", None, False),
    "price": ("Share Price", PRICE, False),
    "market_cap": ("Market Cap ($M)", MILLIONS, True),
    "enterprise_value": ("Enterprise Value ($M)", MILLIONS, True),
    "ev_revenue": ("EV / Revenue", MULTIPLE, False),
    "ev_ebitda": ("EV / EBITDA", MULTIPLE, False),
    "p_fcf": ("P / FCF", MULTIPLE, False),
    "revenue": ("Revenue ($M)", MILLIONS, True),
    "revenue_growth": ("Revenue Growth", PERCENT, False),
    "gross_margin": ("Gross Margin", PERCENT, False),
    "ebitda_margin": ("EBITDA Margin", PERCENT, False),
    "fcf_margin": ("FCF Margin", PERCENT, False),
    "fcf_margin_after_sbc": ("FCF Margin after SBC", PERCENT, False),
    "sbc_pct_revenue": ("SBC % Revenue", PERCENT, False),
    "rule_of_40": ("Rule of 40", PERCENT, False),
    "net_dilution": ("Net Dilution (FY)", PERCENT, False),
    "companies": ("Companies", "0", False),
    "debt": ("+ Debt ($M)", MILLIONS, True),
    "cash": ("- Cash ($M)", MILLIONS, True),
    "short_term_investments": ("- ST Investments ($M)", MILLIONS, True),
    "balance_sheet_date": ("Balance Sheet Date", None, False),
    "debt_basis": ("Debt Basis", None, False),
    "fye_month": ("FYE Month", "0", False),
    "flags": ("Flags", None, False),
}


def _title(ws, title: str, subtitle: str) -> None:
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14, color="1F3864")
    ws["A2"] = subtitle
    ws["A2"].font = Font(italic=True, color="595959")


def _header(ws, row: int, headers: list[str]) -> None:
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 32


def _write_value(cell, value, fmt, scale, label="") -> None:
    if not isinstance(value, str) and pd.isna(value):
        cell.value = label or ""
        cell.font, cell.alignment = LABEL_FONT, Alignment(horizontal="right")
        return
    if isinstance(value, str) and value in BUCKET_NAMES:
        value = BUCKET_NAMES[value]
    cell.value = value / 1e6 if scale else value
    if fmt:
        cell.number_format = fmt


def write_table(ws, df: pd.DataFrame, columns: list[str], start_row: int = 4,
                labels: pd.DataFrame | None = None, stats: pd.DataFrame | None = None,
                index_header: str = "Ticker", widths: dict[str, int] | None = None) -> int:
    """Write df (indexed by ticker or label) with formats; returns the next free row."""
    headers = [index_header] + [COLUMNS.get(c, (c.replace("_", " ").title(), None, False))[0]
                                for c in columns]
    _header(ws, start_row, headers)
    row = start_row + 1
    for key, record in df.iterrows():
        ws.cell(row=row, column=1, value=key).font = Font(bold=True)
        for i, col in enumerate(columns, start=2):
            _, fmt, scale = COLUMNS.get(col, (None, None, False))
            label = labels.at[key, col] if labels is not None and col in labels.columns else ""
            _write_value(ws.cell(row=row, column=i), record[col], fmt, scale, label)
        row += 1

    if stats is not None:
        row += 1
        for stat_name, record in stats.iterrows():
            ws.cell(row=row, column=1, value=stat_name).font = Font(bold=True)
            for i, col in enumerate(columns, start=2):
                cell = ws.cell(row=row, column=i)
                cell.fill = STATS_FILL
                if col in record.index:
                    _, fmt, scale = COLUMNS.get(col, (None, None, False))
                    _write_value(cell, record[col], fmt, scale)
            ws.cell(row=row, column=1).fill = STATS_FILL
            row += 1

    ws.freeze_panes = ws.cell(row=start_row + 1, column=2)
    ws.column_dimensions["A"].width = 18
    for i, col in enumerate(columns, start=2):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(col, 15)
    return row


def _text_sheet(ws, title: str, lines: list[str]) -> None:
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14, color="1F3864")
    for i, line in enumerate(lines, start=3):
        cell = ws.cell(row=i, column=1, value=line)
        if line and not line.startswith(("  ", "•")):
            cell.font = Font(bold=True)
    ws.column_dimensions["A"].width = 130


def write_comps_workbook(path: Path, tables: dict[str, pd.DataFrame], cfg: dict,
                         splits: pd.DataFrame, overrides: pd.DataFrame, market_date: str) -> None:
    comps_cfg = cfg["comps"]
    year = comps_cfg["calendar_year"]
    subtitle = (f"Calendarized to CY{year} · Market data as of {market_date} · "
                "$ in millions except per-share · NM = not meaningful, n/a = missing input")
    comps, labels = tables["comps"].copy(), tables["labels"]
    comps["flags"] = tables["notes"]["notes"].map(lambda notes: " | ".join(notes))

    wb = Workbook()

    ws = wb.active
    ws.title = "Valuation"
    _title(ws, "SaaS Comparable Companies — Valuation", subtitle)
    cols = ["name", "bucket", "price", "market_cap", "enterprise_value",
            "ev_revenue", "ev_ebitda", "p_fcf", "flags"]
    write_table(ws, comps, cols, labels=labels, stats=tables["valuation_stats"],
                widths={"name": 30, "bucket": 24, "flags": 70})

    ws = wb.create_sheet("Operating Metrics")
    _title(ws, "SaaS Comparable Companies — Operating Metrics", subtitle)
    ops = ["name", "bucket", "revenue", "revenue_growth", "gross_margin", "ebitda_margin",
           "fcf_margin", "fcf_margin_after_sbc", "sbc_pct_revenue", "rule_of_40", "net_dilution"]
    write_table(ws, comps, ops, labels=labels, stats=tables["operating_stats"],
                widths={"name": 30, "bucket": 24})

    ws = wb.create_sheet("Profile Medians")
    _title(ws, "Median metrics by company profile",
           "Do the market's multiples line up with growth and profitability?")
    medians = tables["bucket_medians"].rename(index=BUCKET_NAMES)
    write_table(ws, medians, list(medians.columns), index_header="Profile", widths={})
    ws.column_dimensions["A"].width = 28

    ws = wb.create_sheet("EV Bridge")
    _title(ws, "Market cap to enterprise value",
           "EV = market cap + debt - cash - short-term investments. "
           "Balance sheet items are from the latest 10-K.")
    write_table(ws, tables["ev_bridge"],
                ["market_cap", "debt", "cash", "short_term_investments", "enterprise_value",
                 "balance_sheet_date", "debt_basis"], widths={"debt_basis": 48})

    ws = wb.create_sheet("Calendarization")
    _title(ws, f"Fiscal years blended into CY{year}",
           "Each company's fiscal-year figures are weighted by months of overlap with the calendar year.")
    cal = tables["calendarization"]
    write_table(ws, cal, list(cal.columns),
                widths={c: (60 if "method" in c else 16) for c in cal.columns})
    for col in [c for c in cal.columns if c.startswith("revenue")]:
        idx = list(cal.columns).index(col) + 2
        for r in range(5, 5 + len(cal)):
            cell = ws.cell(row=r, column=idx)
            if isinstance(cell.value, (int, float)):
                cell.value, cell.number_format = cell.value / 1e6, MILLIONS

    ws = wb.create_sheet("Stock Splits")
    _title(ws, "Stock splits on record",
           "A split multiplies the share count without changing company value. Share counts "
           "filed before a split are multiplied by its ratio so every year is on today's basis.")
    if splits.empty:
        ws["A4"] = "No stock splits on record."
    else:
        table = (splits.assign(in_data_window=splits["in_data_window"].map({True: "Yes", False: "No"}))
                 .sort_values(["in_data_window", "date"], ascending=[False, False])
                 .set_index("ticker"))
        write_table(ws, table, ["date", "ratio", "source", "share_values_adjusted", "in_data_window"],
                    widths={"source": 40, "share_values_adjusted": 22, "in_data_window": 18})

    ws = wb.create_sheet("Overrides")
    _title(ws, "Manual overrides from 10-K filings",
           "Edit in config/config.yaml (value_usd_millions), then re-run the pipeline and comps.")
    if overrides.empty:
        ws["A4"] = "No overrides configured."
    else:
        write_table(ws, overrides.set_index("ticker"),
                    ["fiscal_year", "concept", "value_usd_millions", "status", "source"],
                    widths={"concept": 26, "source": 95})

    ws = wb.create_sheet("Notes")
    lines = []
    for ticker, record in tables["notes"].iterrows():
        if record["notes"]:
            lines.append(ticker)
            lines += [f"  • {n}" for n in record["notes"]]
            lines.append("")
    _text_sheet(ws, "Per-company notes", lines or ["No caveats."])

    margin_name = {"fcf_margin": "FCF margin", "fcf_margin_after_sbc": "FCF margin after SBC",
                   "ebitda_margin": "EBITDA margin"}[comps_cfg["rule_of_40_margin"]]
    _text_sheet(wb.create_sheet("Methodology"), "Methodology", [
        "Valuation multiples",
        "  EV / Revenue: what the market pays per $1 of sales. The main SaaS multiple, since many SaaS companies have no profits yet.",
        "  EV / EBITDA: price of operating earnings before D&A. NM when EBITDA is negative.",
        "  P / FCF: market cap per $1 of free cash flow, the cash owners could actually take out.",
        f"  Multiples above {comps_cfg['max_meaningful_multiple']}x are shown as NM: they reflect near-zero earnings, not valuation.",
        "",
        "Enterprise value",
        "  EV = market cap + debt - cash"
        + (" - short-term investments" if comps_cfg["ev_net_short_term_investments"] else "")
        + ". It's the price of the whole business, independent of how it's financed.",
        "  Convertible notes count as debt. Operating leases are excluded.",
        "  Simplification: balance sheet is from the latest 10-K, which can be up to a year older than the market price.",
        "",
        "Operating metrics (calendarized)",
        "  EBITDA = operating income + D&A" + (" + stock-based comp" if comps_cfg["ebitda_add_back_sbc"]
                                              else " (GAAP; stock-based comp is NOT added back)") + ".",
        "  FCF = operating cash flow - capex" + (" - capitalized software" if comps_cfg["fcf_subtract_capitalized_software"] else "") + ".",
        "  FCF margin after SBC subtracts stock-based comp too: SBC costs no cash but dilutes shareholders.",
        f"  Rule of 40 = revenue growth + {margin_name}. Above 40% is the common benchmark for a healthy SaaS business.",
        "  Net dilution = change in split-adjusted diluted shares over the last fiscal year.",
        "",
        "Calendarization",
        f"  Fiscal years end in different months, so flow metrics are converted to CY{year}:",
        f"  CY value = (FYE month / 12) x FY ending in {year} + ((12 - FYE month) / 12) x FY ending in {year + 1}.",
        "  If the later fiscal year isn't reported yet, the earlier one is used as a proxy (flagged in Notes).",
        "",
        "Sources",
        "  Financials: SEC EDGAR XBRL companyfacts (10-K filings). Market data: yfinance.",
        "  Every assumption above is set in config/config.yaml.",
    ])

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
